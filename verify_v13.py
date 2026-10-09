# -*- coding: utf-8 -*-
"""Verify v13: status changes, styles, format preservation, analysis appended."""
import openpyxl

path = r"C:\Users\yj202\DoubaoWork\chats\2026-09-07\new-chat\riscv_isa_comparison_v13.xlsx"
wb = openpyxl.load_workbook(path, data_only=False)
ws = wb["RISC-V ISA 全量对比"]
ref = wb["参考文档"]

rows = [27, 65, 66, 85, 97, 98, 135, 137, 143, 147, 156, 157, 158, 173]
print("=== changed rows verification (expect value=-, yellow fill FFFF3CD) ===")
ok = True
for r in rows:
    cell = ws.cell(row=r, column=3)
    fill = cell.fill.fgColor.rgb if cell.fill and cell.fill.fgColor else None
    color = cell.font.color.rgb if cell.font and cell.font.color else None
    status = "OK" if (cell.value == "-" and fill == "00FFF3CD") else "FAIL"
    if status == "FAIL":
        ok = False
    print(f"R{r} {ws.cell(row=r,column=1).value[:30]:32s} C={cell.value!r} fill={fill} font={color} {status}")

print("\n=== untouched cells (expect unchanged) ===")
for r, exp in [(5, "√"), (7, "√"), (16, "-"), (20, "-"), (17, "×"), (23, "√")]:
    v = ws.cell(row=r, column=3).value
    print(f"R{r} {ws.cell(row=r,column=1).value[:30]:32s} C={v!r} expect={exp!r} {'OK' if v==exp else 'FAIL'}")

print("\n=== header/merge preservation ===")
print("A1:", ws["A1"].value)
print("A3:", repr(ws["A3"].value)[:40])
print("C3:", repr(ws["C3"].value)[:40])
print("merged count:", len(list(ws.merged_cells.ranges)))
print("R4 category:", ws["A4"].value)
print("R127 category:", ws["A127"].value)

print("\n=== 参考文档 sheet: original content preserved + analysis appended ===")
print("B1:", ref["B1"].value)
print("B5:", ref["B5"].value)
print("B19:", ref["B19"].value)
found = 0
for r in range(20, 40):
    v = ref.cell(row=r, column=1).value
    if v:
        found += 1
        print(f"R{r}: {v[:70]}")
print("appended lines:", found)
