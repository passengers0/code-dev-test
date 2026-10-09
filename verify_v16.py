# -*- coding: utf-8 -*-
"""Verify v16 color coding."""
import openpyxl

path = r"C:\Users\yj202\DoubaoWork\chats\2026-09-07\new-chat\riscv_isa_comparison_v16.xlsx"
wb = openpyxl.load_workbook(path, data_only=False)
ns = wb["汽车标准不支持组合及原因"]

def fill(row, col):
    c = ns.cell(row=row, column=col)
    return c.fill.fgColor.rgb if c.fill and c.fill.fgColor else None

expect_rows = {3:"FFF4CCCC",4:"FFF4CCCC",5:"FFF4CCCC",6:"FFD6E4F0",7:"FFD6E4F0",
               8:"FFFCE4D6",9:"FFFCE4D6",10:"FFD9EAD3",11:"FFD9EAD3",12:"FFD6E4F0",13:"FFE1D5E7"}
ok = True
for r, exp in expect_rows.items():
    got = fill(r, 2)
    good = got == exp
    ok = ok and good
    print(f"R{r} {ns.cell(row=r,column=2).value[:20]:22s} fill={got} expect={exp} {'OK' if good else 'FAIL'}")

print("\nlegend rows:")
for r in range(14, 19):
    print(f"  R{r}: A fill={fill(r,1)} | B={str(ns.cell(row=r,column=2).value)[:50]}")
print("note R19:", str(ns.cell(row=19, column=1).value)[:50], "...")
print("merged:", sorted(str(m) for m in ns.merged_cells.ranges))
print("ALL OK:", ok)

# main sheet untouched
ws = wb["RISC-V ISA 全量对比"]
print("\nmain sheet P R17 C:", ws["C17"].value, "| Zks R85:", ws["C85"].value, "| merged:", len(list(ws.merged_cells.ranges)))
