# -*- coding: utf-8 -*-
"""Verify v14."""
import openpyxl

path = r"C:\Users\yj202\DoubaoWork\chats\2026-09-07\new-chat\riscv_isa_comparison_v14.xlsx"
wb = openpyxl.load_workbook(path, data_only=False)
ws = wb["RISC-V ISA 全量对比"]
ref = wb["参考文档"]

def info(r):
    c = ws.cell(row=r, column=3)
    fill = c.fill.fgColor.rgb if c.fill and c.fill.fgColor else None
    return c.value, fill

expect = {
    85: ("√", "FFC6EFCE"),      # Zks green
    97: ("-", "FFFFF3CD"),      # Zvkng yellow
    98: ("-", "FFFFF3CD"),      # Zvksg yellow
    27: ("×", "FFFFC7CE"),      # Zacas red
    65: ("×", "FFFFC7CE"),
    66: ("×", "FFFFC7CE"),
    135: ("×", "FFFFC7CE"),
    137: ("×", "FFFFC7CE"),
    143: ("×", "FFFFC7CE"),
    147: ("×", "FFFFC7CE"),
    156: ("×", "FFFFC7CE"),
    157: ("×", "FFFFC7CE"),
    158: ("×", "FFFFC7CE"),
    173: ("×", "FFFFC7CE"),
}
allok = True
for r, (ev, ef) in expect.items():
    v, f = info(r)
    ok = (v == ev and f == ef)
    allok = allok and ok
    print(f"R{r} {ws.cell(row=r,column=1).value[:28]:30s} C={v!r} fill={f} expect={ev!r}/{ef} {'OK' if ok else 'FAIL'}")

# spot check untouched cells
for r, exp in [(5, "√"), (20, "-"), (16, "-"), (29, "√"), (128, "√"), (130, "-")]:
    v = ws.cell(row=r, column=3).value
    print(f"R{r} {ws.cell(row=r,column=1).value[:24]:26s} C={v!r} expect={exp!r} {'OK' if v==exp else 'FAIL'}")

print("\nheader A1:", ws["A1"].value[:40])
print("merged count:", len(list(ws.merged_cells.ranges)))
print("\n参考文档 analysis head:")
for r in range(21, 25):
    print("R%d: %s" % (r, (ref.cell(row=r, column=1).value or "")[:80]))
print("...")
for r in range(34, 37):
    print("R%d: %s" % (r, (ref.cell(row=r, column=1).value or "")[:80]))
print("\nALL OK:", allok)
