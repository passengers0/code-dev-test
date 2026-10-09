# -*- coding: utf-8 -*-
"""Verify v15."""
import openpyxl

path = r"C:\Users\yj202\DoubaoWork\chats\2026-09-07\new-chat\riscv_isa_comparison_v15.xlsx"
wb = openpyxl.load_workbook(path, data_only=False)
ws = wb["RISC-V ISA 全量对比"]
ref = wb["参考文档"]
ns = wb["汽车标准不支持组合及原因"]

print("sheets:", wb.sheetnames)

# 1. P row
c = ws["C17"]
print(f"\nP R17: C={c.value!r} fill={c.fill.fgColor.rgb} expect '-'/FFFFF3CD ->",
      "OK" if c.value == "-" and c.fill.fgColor.rgb == "FFFFF3CD" else "FAIL")
print("note A186:", str(ws["A186"].value)[:60], "...")

# 2. 参考文档 restored
vals = [ref.cell(row=r, column=1).value for r in range(20, 38)]
print("\n参考文档 R20-R37 all empty:", all(v is None for v in vals), "| R1:", ref["A1"].value, "| R19:", ref["A19"].value)

# 3. new sheet content
print(f"\n新sheet dims: {ns.dimensions}")
print("A1:", ns["A1"].value)
hdr = [ns.cell(row=2, column=c).value for c in range(1, 9)]
print("headers:", hdr)
count = 0
for r in range(3, 30):
    a = ns.cell(row=r, column=1).value
    if a is None:
        break
    count += 1
    print(f"  row{r}: [{a}] {ns.cell(row=r,column=2).value} | {str(ns.cell(row=r,column=6).value)[:30]}")
print("data rows:", count, "(expect 11)")

# 4. main sheet spot checks
for r, exp in [(85, "√"), (97, "-"), (98, "-"), (27, "×"), (65, "×"), (173, "×"), (5, "√"), (20, "-")]:
    v = ws.cell(row=r, column=3).value
    if v != exp:
        print(f"  MISMATCH R{r}: {v!r} != {exp!r}")
print("main sheet checks done, merged:", len(list(ws.merged_cells.ranges)))
