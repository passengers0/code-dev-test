# -*- coding: utf-8 -*-
"""Identify unsupported-by-HighTec automotive extensions and inspect styles."""
import openpyxl

path = r"C:\Users\yj202\DoubaoWork\chats\2026-09-07\new-chat\riscv_isa_comparison_v12.xlsx"
wb = openpyxl.load_workbook(path, data_only=False)
ws = wb["RISC-V ISA 全量对比"]

# candidate rows: B(草案) in {√,-} and C(HighTec) == ×
cands = []
for r in range(5, 177):
    a = ws.cell(row=r, column=1).value
    b = ws.cell(row=r, column=2).value
    c = ws.cell(row=r, column=3).value
    if a and b in ("√", "-") and c == "×":
        cands.append((r, a, b, c))
print("candidates (draft √/-, HighTec ×):", len(cands))
for r, a, b, c in cands:
    print(r, a, "| draft:", b, "| hightec:", c)

# inspect styles of C column symbols on a few rows
def style_of(coord):
    cell = ws[coord]
    f = cell.font
    fill = cell.fill
    fg = fill.fgColor.rgb if fill and fill.fgColor else None
    return f"font_color={f.color.rgb if f and f.color else None} bold={f.bold} fill={fg} fill_type={fill.patternType if fill else None}"

print("\n--- C col styles ---")
for coord in ["C5", "C7", "C16", "C20", "C27", "C42"]:  # √,-,×,-,×,-
    print(coord, ws[coord].value, style_of(coord))
print("\n--- B col styles (draft) sample ---")
for coord in ["B27", "B65", "B17", "B5"]:
    print(coord, ws[coord].value, style_of(coord))
