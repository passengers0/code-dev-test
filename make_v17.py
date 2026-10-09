# -*- coding: utf-8 -*-
"""Generate v17: two-category color coding in '汽车标准不支持组合及原因'.
Category A (requires compiler support, light red): Zacas, Zicfilp, Zicfiss
Category B (non-compiler support, light blue): the other 8
Legend simplified to 2 rows; note moves to R17."""
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

SRC = r"C:\Users\yj202\DoubaoWork\chats\2026-09-07\new-chat\riscv_isa_comparison_v15.xlsx"
OUT = r"C:\Users\yj202\DoubaoWork\chats\2026-09-07\new-chat\riscv_isa_comparison_v17.xlsx"

wb = openpyxl.load_workbook(SRC)
ns = wb["汽车标准不支持组合及原因"]

thin = Side(style="thin", color="FFBFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(vertical="top", wrap_text=True)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

# 1. free R15 note
if "A15:H15" in [str(m) for m in ns.merged_cells.ranges]:
    ns.unmerge_cells("A15:H15")
note_text = ns.cell(row=15, column=1).value
for c in range(1, 9):
    ns.cell(row=15, column=c).value = None

# 2. two-category coloring
RED = "FFF4CCCC"    # requires compiler support
BLUE = "FFD6E4F0"   # non-compiler support
CAT = {
    3:  RED,   # Zacas
    4:  RED,   # Zicfilp
    5:  RED,   # Zicfiss
    6:  BLUE,  # Smdbltrp
    7:  BLUE,  # Smcntrpmf
    8:  BLUE,  # Sspmp
    9:  BLUE,  # Sspm
    10: BLUE,  # Sv48
    11: BLUE,  # Sv57
    12: BLUE,  # Svvptc
    13: BLUE,  # Sdtrig
}
for r, fill in CAT.items():
    fillobj = PatternFill("solid", fgColor=fill)
    for c in range(1, 9):
        cell = ns.cell(row=r, column=c)
        cell.fill = fillobj
        cell.border = border
        cell.alignment = wrap if c >= 5 else center
    ns.cell(row=r, column=4).font = Font(color="FF9C0006", bold=True)
    print(f"R{r} {ns.cell(row=r,column=2).value[:22]:24s} -> {'要求编译器支持' if fill==RED else '非编译器支持'}")

# 3. simplified legend (2 rows at R14-15)
legend = [
    (RED,  "要求编译器支持（需 LLVM/Clang 实现指令选择或 CFI 全链路插桩，库只能模拟或汇编补丁）：Zacas / Zicfilp / Zicfiss"),
    (BLUE, "非编译器支持（硬件/固件/OS/调试器范畴，编译器无指令需求或不涉及，由硬件 CSR、OS 配置或调试工具链实现）：Smdbltrp / Smcntrpmf / Sspmp / Sspm / Sv48 / Sv57 / Svvptc / Sdtrig"),
]
row = 14
for fill, text in legend:
    a = ns.cell(row=row, column=1)
    a.fill = PatternFill("solid", fgColor=fill)
    a.border = border
    b = ns.cell(row=row, column=2, value=text)
    b.alignment = Alignment(vertical="center", wrap_text=True)
    b.font = Font(size=9)
    ns.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
    for c in range(2, 9):
        ns.cell(row=row, column=c).border = border
    ns.row_dimensions[row].height = 30
    row += 1

# 4. note to R17 (R16 blank)
t = ns.cell(row=17, column=1, value=note_text)
t.font = Font(italic=True, size=9, color="FF595959")
t.alignment = wrap
ns.merge_cells("A17:H17")
ns.row_dimensions[17].height = 30

wb.save(OUT)
print("saved:", OUT)
