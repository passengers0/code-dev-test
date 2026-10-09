# -*- coding: utf-8 -*-
"""Dump all data rows of v12 comparison sheet."""
import openpyxl

path = r"C:\Users\yj202\DoubaoWork\chats\2026-09-07\new-chat\riscv_isa_comparison_v12.xlsx"
wb = openpyxl.load_workbook(path, data_only=False)
ws = wb["RISC-V ISA 全量对比"]
out = []
for r in range(1, ws.max_row + 1):
    a = ws.cell(row=r, column=1).value
    b = ws.cell(row=r, column=2).value
    c = ws.cell(row=r, column=3).value
    if a is None and b is None and c is None:
        continue
    out.append(f"R{r}\tA={a!r}\tB={b!r}\tC={c!r}")
open(r"C:\work\code\riscv11_test\v12_dump.txt", "w", encoding="utf-8").write("\n".join(out))
print("lines:", len(out))
